import os
import io
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.services.document_parser import DocumentParser, MAX_FILE_SIZE_BYTES
from app.store.state_store import store
import docx

client = TestClient(app)

def test_parse_txt():
    sample_text = (
        "Facility: St. Jude Hospital\n"
        "Product: Epinephrine Injection 1mg/mL, USP\n"
        "Batch: B24017\n"
        "Quantity Affected: 50 vials\n"
        "Defect: Rubber stopper leakage discovered during unpacking."
    )
    text, fmt = DocumentParser.parse_file("complaint.txt", sample_text.encode("utf-8"))
    assert fmt == "TXT"
    assert "Epinephrine" in text

    entities = DocumentParser.extract_complaint_entities(text)
    assert entities["product_name"] == "Epinephrine Injection"
    assert entities["batch_number"] == "B24017"
    assert entities["affected_quantity"] == 50
    assert entities["complaint_type"] == "Container Closure Leakage"

def test_parse_docx():
    doc = docx.Document()
    doc.add_heading("Customer Quality Incident Notice", 0)
    doc.add_paragraph("Product: Metformin Hydrochloride")
    doc.add_paragraph("Batch Number: M9012")
    doc.add_paragraph("Quantity: 120 bottles")
    doc.add_paragraph("Issue: Bottles missing child-safety closures.")

    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    text, fmt = DocumentParser.parse_file("incident.docx", content)
    assert fmt == "DOCX"
    assert "Metformin" in text

    entities = DocumentParser.extract_complaint_entities(text)
    assert entities["product_name"] == "Metformin Hydrochloride"
    assert entities["batch_number"] == "M9012"
    assert entities["affected_quantity"] == 120
    assert "Packaging" in entities["complaint_type"]

def test_parse_eml():
    eml_content = (
        "From: pharmacovigilance@mercyhealth.org\n"
        "To: complaints@pharmaco.com\n"
        "Subject: Urgent: Defective Saline Infusion Lot NS4401\n"
        "Date: Wed, 11 Mar 2026 10:30:00 -0500\n"
        "Content-Type: text/plain; charset=utf-8\n\n"
        "Dear Quality Team,\n\n"
        "We received Normal Saline 0.9% infusion bags (Lot NS4401).\n"
        "Staff observed floating black specks in 15 bags.\n"
        "Please initiate immediate investigation.\n"
    ).encode("utf-8")

    text, fmt = DocumentParser.parse_file("complaint_email.eml", eml_content)
    assert fmt == "EML"
    assert "mercyhealth.org" in text

    entities = DocumentParser.extract_complaint_entities(text)
    assert "Saline" in entities["product_name"]
    assert entities["batch_number"] == "NS4401"
    assert entities["affected_quantity"] == 15
    assert entities["complaint_type"] == "Foreign Particulate Matter"

def test_parse_real_pdf_file():
    pdf_path = os.path.join(os.path.dirname(__file__), "..", "sample_data", "sample_complaint_letter.pdf")
    if os.path.exists(pdf_path):
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        text, fmt = DocumentParser.parse_file("sample_complaint_letter.pdf", pdf_bytes)
        assert fmt == "PDF"
        assert "Epinephrine" in text
        entities = DocumentParser.extract_complaint_entities(text)
        assert entities["product_name"] == "Epinephrine Injection"
        assert entities["batch_number"] == "B24017"

def test_unsupported_format_and_oversize():
    with pytest.raises(ValueError, match="Unsupported format"):
        DocumentParser.parse_file("archive.zip", b"PK12345")

    oversize_bytes = b"0" * (MAX_FILE_SIZE_BYTES + 1024)
    with pytest.raises(ValueError, match="exceeds maximum limit"):
        DocumentParser.parse_file("large.pdf", oversize_bytes)

def test_api_upload_endpoint():
    session_id = "test-doc-upload-session"
    store.reset_session(session_id)

    file_content = (
        "Customer: St. Jude Hospital Pharmacy\n"
        "Product: Epinephrine Injection 1mg/mL\n"
        "Batch: B24017\n"
        "Quantity: 50 vials\n"
        "Description: Vials leaking around rubber seal."
    ).encode("utf-8")

    response = client.post(
        "/api/documents/upload",
        data={"session_id": session_id},
        files={"file": ("incident_report.txt", file_content, "text/plain")}
    )

    assert response.status_code == 200
    data = response.json()
    assert data["format"] == "TXT"
    assert data["filename"] == "incident_report.txt"
    assert data["extracted_fields_count"] >= 4
    assert any("Document Extraction" in b for b in data["tool_badges"])

    complaint = data["complaint"]
    assert complaint["product_batch"]["product_name"] == "Epinephrine Injection"
    assert complaint["manufacturing"]["batch_number"] == "B24017"
    assert complaint["manufacturing"]["affected_quantity"] == 50
    assert complaint["risk_assessment"]["severity"] == "CRITICAL"

    # CRITICAL: Verify follow-up natural language edit after document extraction!
    edit_response = client.post("/api/chat", json={
        "session_id": session_id,
        "message": "The affected quantity was actually 75 units, and customer reported it happened across 3 boxes."
    })
    assert edit_response.status_code == 200
    edit_data = edit_response.json()
    updated_complaint = edit_data["complaint"]
    assert updated_complaint["manufacturing"]["affected_quantity"] == 75
    # Preserved fields from document upload:
    assert updated_complaint["product_batch"]["product_name"] == "Epinephrine Injection"
    assert updated_complaint["manufacturing"]["batch_number"] == "B24017"

def test_api_paste_endpoint():
    session_id = "test-paste-session"
    store.reset_session(session_id)

    text = "Product: Metformin 500mg\nBatch: M880\nQuantity: 30 bottles\nIssue: Damaged packaging."
    response = client.post("/api/documents/paste", json={
        "session_id": session_id,
        "text": text
    })

    assert response.status_code == 200
    data = response.json()
    assert data["complaint"]["manufacturing"]["batch_number"] == "M880"
    assert data["complaint"]["manufacturing"]["affected_quantity"] == 30

def test_paste_flow_cardiosafe_regression():
    """
    Regression test verifying the complete Paste Complaint Text / Email flow
    with the exact CardioSafe 10 mg Tablets quality complaint.
    Validates all 16 expected structured fields and full description preservation.
    """
    session_id = "test-cardiosafe-paste-regression"
    store.reset_session(session_id)

    raw_complaint_text = (
        "Subject: Critical Quality Complaint – CardioSafe 10 mg Tablets, Batch CS240817\n\n"
        "Dear Quality Assurance Team,\n\n"
        "I am writing to formally report a quality complaint regarding CardioSafe 10 mg Tablets, "
        "manufactured by ApexPharma Laboratories. The product is supplied as 10 mg immediate-release "
        "tablets in PVC-Alu blister packaging.\n\n"
        "The complaint was received from St. Jude Memorial Hospital on 13 September 2026. "
        "The affected product is from Batch Number CS240817, manufactured on 17 August 2026, "
        "with an expiration date of 16 August 2028.\n\n"
        "The hospital reports that 25 tablets were affected. Upon opening an intact blister pack, "
        "several tablets were found to be severely chipped and cracked, with visible powder present "
        "inside the blister pockets. The affected tablets also showed abnormal yellow discoloration "
        "compared with the standard product appearance.\n\n"
        "The complaint is classified as a Product Quality / Physical Defect involving tablet damage, "
        "powder leakage, and abnormal discoloration. The product was purchased from Cornerstone "
        "Pharmacy, Downtown, under Rx Number 9982415.\n\n"
        "The affected samples, original blister packaging, and purchase receipt have been securely "
        "retained and are available for laboratory investigation. Please initiate the appropriate "
        "QA investigation and batch evaluation and provide recommendations regarding further "
        "handling of the affected product."
    )

    response = client.post("/api/documents/paste", json={
        "session_id": session_id,
        "text": raw_complaint_text
    })

    assert response.status_code == 200
    data = response.json()
    complaint = data["complaint"]

    # Verify origin & customer details
    assert complaint["origin_customer"]["complaint_source"] == "Hospital / Healthcare Facility"
    assert complaint["origin_customer"]["customer_name"] == "St. Jude Memorial Hospital"
    assert complaint["origin_customer"]["complaint_type"] == "Product Quality / Physical Defect"
    assert complaint["origin_customer"]["complaint_date"] == "13 September 2026"
    assert complaint["origin_customer"]["purchase_location"] == "Cornerstone Pharmacy, Downtown"
    assert complaint["origin_customer"]["rx_number"] == "9982415"

    # Verify product & batch details
    assert complaint["product_batch"]["product_name"] == "CardioSafe"
    assert complaint["product_batch"]["product_grade_strength"] == "10 mg"
    assert complaint["product_batch"]["dosage_form"] == "Immediate-release tablets"
    assert complaint["product_batch"]["packaging"] == "PVC-Alu blister"
    assert complaint["product_batch"]["manufacturer"] == "ApexPharma Laboratories"

    # Verify manufacturing details
    assert complaint["manufacturing"]["batch_number"] == "CS240817"
    assert complaint["manufacturing"]["manufacturing_date"] == "17 August 2026"
    assert complaint["manufacturing"]["expiry_date"] == "16 August 2028"
    assert complaint["manufacturing"]["affected_quantity"] == 25
    assert complaint["manufacturing"]["unit_of_measure"] == "Tablets"

    # Verify full complaint narrative is preserved without truncation
    desc = complaint["origin_customer"]["complaint_description"]
    assert "CardioSafe 10 mg Tablets" in desc
    assert "St. Jude Memorial Hospital" in desc
    assert "Cornerstone Pharmacy, Downtown" in desc
    assert "laboratory investigation" in desc
    assert len(desc) > 300

    # Verify risk assessment was triggered and populated under ICH Q9
    assert complaint["risk_assessment"]["severity"] in ["CRITICAL", "HIGH", "MEDIUM"]
    assert len(complaint["risk_assessment"]["recommended_actions"]) > 0

def test_upload_flow_cardiosafe_pdf_regression():
    """
    Regression test verifying drag & drop PDF upload flow using sample_data/CardioSafe_Quality_Complaint.pdf.
    Verifies that all 16 pharmaceutical fields and full narrative description are correctly extracted
    and identical to the paste flow extraction without mapping narrative fragments to customer name.
    """
    session_id = "test-cardiosafe-pdf-regression"
    store.reset_session(session_id)

    pdf_path = os.path.join(os.path.dirname(__file__), "..", "..", "sample_data", "CardioSafe_Quality_Complaint.pdf")
    if not os.path.exists(pdf_path):
        pdf_path = os.path.join(os.path.dirname(__file__), "..", "sample_data", "CardioSafe_Quality_Complaint.pdf")

    assert os.path.exists(pdf_path), f"CardioSafe PDF not found at {pdf_path}"

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    response = client.post(
        "/api/documents/upload",
        data={"session_id": session_id},
        files={"file": ("CardioSafe_Quality_Complaint.pdf", pdf_bytes, "application/pdf")}
    )

    assert response.status_code == 200, f"Upload failed: {response.text}"
    data = response.json()
    complaint = data["complaint"]

    # Verify origin & customer details - NO narrative fragments mapped to customer name
    assert complaint["origin_customer"]["complaint_source"] == "Hospital / Healthcare Facility"
    assert complaint["origin_customer"]["customer_name"] == "St. Jude Memorial Hospital"
    assert complaint["origin_customer"]["complaint_type"] == "Product Quality / Physical Defect"
    assert complaint["origin_customer"]["complaint_date"] == "13 September 2026"
    assert complaint["origin_customer"]["purchase_location"] == "Cornerstone Pharmacy, Downtown"
    assert complaint["origin_customer"]["rx_number"] == "9982415"

    # Verify product & batch details
    assert complaint["product_batch"]["product_name"] == "CardioSafe"
    assert complaint["product_batch"]["product_grade_strength"] == "10 mg"
    assert complaint["product_batch"]["dosage_form"] == "Immediate-release tablets"
    assert complaint["product_batch"]["packaging"] == "PVC-Alu blister"
    assert complaint["product_batch"]["manufacturer"] == "ApexPharma Laboratories"

    # Verify manufacturing details
    assert complaint["manufacturing"]["batch_number"] == "CS240817"
    assert complaint["manufacturing"]["manufacturing_date"] == "17 August 2026"
    assert complaint["manufacturing"]["expiry_date"] == "16 August 2028"
    assert complaint["manufacturing"]["affected_quantity"] == 25
    assert complaint["manufacturing"]["unit_of_measure"] == "Tablets"

    # Verify narrative description preserved without truncation
    desc = complaint["origin_customer"]["complaint_description"]
    assert "CardioSafe 10 mg Tablets" in desc
    assert "St. Jude Memorial Hospital" in desc
    assert "Cornerstone Pharmacy, Downtown" in desc
    assert "laboratory investigation" in desc
    assert len(desc) > 300

    # Verify risk assessment
    assert complaint["risk_assessment"]["severity"] in ["CRITICAL", "HIGH", "MEDIUM"]
    assert len(complaint["risk_assessment"]["recommended_actions"]) > 0

